"""Bounded document extraction and conservative specification matching."""
import csv
import io
import os
import re
import shutil
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor

from ekt_client import ekt_client, number
from settings import ServiceError

MAX_BYTES = 15 * 1024 * 1024
MAX_ROWS = 250
MAX_CHARS = 100000
SUPPORTED = {".pdf", ".txt", ".docx", ".xlsx", ".xls", ".jpg", ".jpeg", ".png", ".webp"}


def ocr_available():
    executable = os.getenv("TESSERACT_CMD") or shutil.which("tesseract")
    return bool(executable and os.path.isfile(executable))


class SpecificationParser:
    def __init__(self, catalog=ekt_client):
        self.catalog = catalog

    @staticmethod
    def validate(content, filename):
        extension = os.path.splitext(filename.lower())[1]
        if extension not in SUPPORTED:
            raise ServiceError("Поддерживаются PDF, TXT, DOCX, XLSX, XLS, JPEG, PNG, WEBP. Файл .doc сохраните как .docx.", "unsupported_file", 415)
        if not content:
            raise ServiceError("Файл пуст.", "empty_file", 422)
        if len(content) > MAX_BYTES:
            raise ServiceError("Файл превышает 15 МиБ.", "file_too_large", 413)
        signatures = {".pdf": b"%PDF-", ".xls": b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1",
                      ".png": b"\x89PNG\r\n\x1a\n", ".jpg": b"\xff\xd8\xff", ".jpeg": b"\xff\xd8\xff"}
        if extension in signatures and not content.startswith(signatures[extension]):
            raise ServiceError("Содержимое файла не соответствует расширению.", "file_signature_mismatch", 415)
        if extension == ".webp" and not (content[:4] == b"RIFF" and content[8:12] == b"WEBP"):
            raise ServiceError("Повреждённый WEBP.", "file_signature_mismatch", 415)
        if extension in {".docx", ".xlsx"}:
            try:
                with zipfile.ZipFile(io.BytesIO(content)) as archive:
                    entries = archive.infolist()
                    marker = "word/document.xml" if extension == ".docx" else "xl/workbook.xml"
                    if marker not in archive.namelist():
                        raise ValueError("missing_document")
                    if len(entries) > 2000 or sum(entry.file_size for entry in entries) > 50 * 1024 * 1024:
                        raise ServiceError("Распакованное содержимое слишком велико.", "archive_too_large", 413)
                    if any(entry.flag_bits & 1 for entry in entries):
                        raise ServiceError("Защищённые паролем файлы не поддерживаются.", "encrypted_document", 422)
            except (zipfile.BadZipFile, ValueError):
                raise ServiceError("Файл не является корректным DOCX/XLSX.", "file_signature_mismatch", 415)
        return extension

    @staticmethod
    def _ocr(image):
        if not ocr_available():
            raise ServiceError("Для фото и сканов установите Tesseract OCR с языками rus, eng и kaz. Либо отправьте текстовую спецификацию.", "ocr_unavailable", 503)
        import pytesseract
        if os.getenv("TESSERACT_CMD"):
            pytesseract.pytesseract.tesseract_cmd = os.environ["TESSERACT_CMD"]
        languages = set(pytesseract.get_languages(config=""))
        chosen = [lang for lang in ("rus", "eng", "kaz") if lang in languages]
        if not chosen:
            raise ServiceError("В Tesseract отсутствуют языковые данные rus/eng/kaz.", "ocr_languages_missing", 503)
        try:
            return pytesseract.image_to_string(image, lang="+".join(chosen), timeout=10)
        except RuntimeError:
            raise ServiceError("Распознавание заняло слишком много времени. Отправьте меньший фрагмент.", "ocr_timeout", 422)

    def extract_text(self, content, filename):
        extension = self.validate(content, filename)
        if extension == ".txt":
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                raise ServiceError("TXT должен быть в UTF-8.", "invalid_text_encoding", 422)
            if "\x00" in text:
                raise ServiceError("В TXT обнаружены бинарные данные.", "file_signature_mismatch", 415)
            return text
        if extension == ".pdf":
            import pymupdf
            from PIL import Image
            with pymupdf.open(stream=content, filetype="pdf") as document:
                if document.needs_pass:
                    raise ServiceError("PDF защищён паролем.", "encrypted_document", 422)
                if len(document) > 50:
                    raise ServiceError("В PDF больше 50 страниц.", "too_many_pages", 413)
                text, scans = [], 0
                for page in document:
                    words = page.get_text("words", sort=True)
                    if words:
                        rows = []
                        for word in sorted(words, key=lambda w: (round(w[1] / 3), w[0])):
                            if not rows or abs(word[1] - rows[-1][0]) > 3:
                                rows.append((word[1], [word]))
                            else:
                                rows[-1][1].append(word)
                        for _, row in rows:
                            row.sort(key=lambda w: w[0])
                            line, last_x = "", None
                            for word in row:
                                line += (("\t" if word[0] - last_x > 15 else " ") if last_x is not None else "") + word[4]
                                last_x = word[2]
                            text.append(line)
                    else:
                        scans += 1
                        if scans > 10:
                            raise ServiceError("Разделите скан на файлы не более 10 страниц.", "too_many_scans", 413)
                        if page.rect.width * page.rect.height * 2.25 > 16000000:
                            raise ServiceError("Страница скана слишком велика.", "image_too_large", 413)
                        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False)
                        text.append(self._ocr(Image.open(io.BytesIO(pixmap.tobytes("png")))))
                    if sum(map(len, text)) > MAX_CHARS:
                        raise ServiceError("В документе слишком много текста.", "text_too_large", 413)
                return "\n".join(text)
        if extension == ".docx":
            from docx import Document
            doc = Document(io.BytesIO(content))
            lines = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                lines.extend("\t".join(cell.text.strip() for cell in row.cells) for row in table.rows)
            return "\n".join(lines)
        if extension == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
            try:
                lines = []
                for sheet in workbook:
                    if sheet.max_row and sheet.max_row > 2000 or sheet.max_column and sheet.max_column > 100:
                        raise ServiceError("Таблица слишком велика.", "table_too_large", 413)
                    for index, row in enumerate(sheet.iter_rows(values_only=True)):
                        if index >= 2000:
                            raise ServiceError("Таблица слишком велика.", "table_too_large", 413)
                        lines.append("\t".join("" if value is None else str(value) for value in row))
                        if sum(map(len, lines)) > MAX_CHARS:
                            raise ServiceError("В таблице слишком много текста.", "text_too_large", 413)
                return "\n".join(lines)
            finally:
                workbook.close()
        if extension == ".xls":
            import xlrd
            workbook = xlrd.open_workbook(file_contents=content)
            lines = []
            for sheet in workbook.sheets():
                if sheet.nrows > 2000 or sheet.ncols > 100:
                    raise ServiceError("Таблица слишком велика.", "table_too_large", 413)
                lines.extend("\t".join(map(str, sheet.row_values(index))) for index in range(sheet.nrows))
            return "\n".join(lines)
        from PIL import Image
        with Image.open(io.BytesIO(content)) as image:
            if image.width * image.height > 16000000:
                raise ServiceError("Изображение превышает 16 мегапикселей.", "image_too_large", 413)
            image.load()
            return self._ocr(image)

    @staticmethod
    def rows(text):
        if len(text) > MAX_CHARS:
            raise ServiceError("Документ превышает лимит извлечённого текста.", "text_too_large", 413)
        header, output, ignored = None, [], []
        for line in text.splitlines():
            if not line.strip():
                continue
            delimiter = "\t" if "\t" in line else ";" if ";" in line else None
            try:
                cells = [cell.strip() for cell in next(csv.reader([line], delimiter=delimiter, strict=True))] if delimiter else [line.strip()]
            except csv.Error:
                raise ServiceError("В табличной строке нарушены кавычки или разделители. Проверьте файл.", "invalid_table_row", 422)
            lowered = [cell.lower() for cell in cells]
            quantity_columns = [i for i, cell in enumerate(lowered) if re.fullmatch(r"кол[ -]?во\.?|количество|qty|quantity|саны", cell)]
            if quantity_columns:
                header = {"quantity": quantity_columns[0], "delimiter": delimiter,
                          "article": next((i for i, cell in enumerate(lowered) if cell in {"артикул", "код", "sku", "article"}), None),
                          "name": next((i for i, cell in enumerate(lowered) if cell in {"наименование", "название", "description", "товар"}), None),
                          "unit": next((i for i, cell in enumerate(lowered) if re.fullmatch(r"ед\.?\s*изм\.?|единица|unit", cell)), None)}
                ignored.append(line)
                continue
            if re.match(r"^(?:спецификация|specification|проект|заказчик|дата|объект|подпись|страница|итого|наименование)\b", line.lower()):
                ignored.append(line)
                continue
            quantity, unit, query = None, None, line
            if header and delimiter and delimiter == header["delimiter"] and len(cells) > header["quantity"]:
                raw_quantity = cells[header["quantity"]]
                # A quantity column is explicit. Preserve integer-valued Excel
                # cells ("3.0"), but do not coerce exponents, signs or decimals.
                quantity = number(raw_quantity) if re.fullmatch(r"\d+(?:[.,]0+)?", raw_quantity) else None
                selected = [cells[header[key]] for key in ("article", "name") if header[key] is not None and header[key] < len(cells)]
                query = " ".join(selected) if selected else " ".join(cells[:header["quantity"]])
                if header["unit"] is not None and header["unit"] < len(cells):
                    unit = cells[header["unit"]]
            else:
                match = re.search(r"(?<![\w.,-])(\d+(?:[.,]\d+)?)\s*(шт\.?|штук|ед\.?|дана|pcs|м)(?![\w²])", line, re.I)
                if not match:
                    match = re.search(r"(?:шт\.?|дана|pcs|ед\.?)\s+(\d+(?:[.,]\d+)?)\s*$", line, re.I)
                if match:
                    quantity = number(match[1])
                    unit = "м" if re.search(r"\d\s*м\b", match[0]) else "шт."
                    query = line[:match.start()] + " " + line[match.end():]
            valid_quantity = quantity is not None and quantity > 0 and quantity.is_integer() and quantity <= 100000
            output.append({"query_line": line, "query": query.strip(), "quantity": int(quantity) if valid_quantity else None,
                           "unit": unit, "quantity_warning": None if valid_quantity else "Уточните целое положительное количество."})
        if len(output) > MAX_ROWS:
            raise ServiceError("В спецификации больше 250 позиций. Разделите файл.", "too_many_rows", 413)
        return output, ignored

    def _match_row(self, row):
        results = self.catalog.search_products(row["query"], limit=3)
        exact = [item for item in results if item.get("_match_type") == "exact"]
        candidates = exact or [item for item in results if item.get("_match_score", 0) >= 2 and item.get("_match_coverage", 0) >= 0.65]
        if not candidates or len(candidates) > 1 and candidates[0].get("_match_score", 0) == candidates[1].get("_match_score", 0):
            return None, {**row, "status": "Нет однозначного совпадения; уточните артикул.", "candidates": [{"id": item["id"], "name": item.get("name")} for item in candidates]}
        target = candidates[0]
        detail = self.catalog.get_product_detail(target["id"])
        product = detail or target
        price = product.get("price") if detail and detail.get("price_verified") else None
        stock = product.get("quantity") if detail and detail.get("stock_verified") else None
        quantity = row["quantity"]
        subtotal = round(price * quantity, 2) if price is not None and quantity is not None else None
        analogs = self.catalog.find_analogs(detail, limit=1) if stock == 0 else []
        status = ("Количество требует уточнения" if quantity is None else "Остаток не проверен" if stock is None
                  else "В наличии" if stock >= quantity else "Частично в наличии" if stock > 0 else "Нет в наличии")
        return {**row, "product_id": product["id"], "name": product["name"], "article": product.get("article", ""),
                "unit_price": price, "subtotal": subtotal, "stock_available": stock, "stock_verified": stock is not None,
                "price_verified": price is not None, "status": status, "unit": row["unit"] or product.get("unit"),
                "image": product.get("image"), "url": product.get("url"),
                "specifications": product.get("specifications", []), "certificate_url": product.get("certificate_url"),
                "last_checked_at": detail.get("last_checked_at") if detail else None,
                "data_quality_warnings": product.get("data_quality_warnings", []),
                "analog": analogs[0] if analogs else None}, None

    def parse_specification(self, content, filename="specification.pdf"):
        try:
            text = self.extract_text(content, filename)
        except ServiceError:
            raise
        except Exception:
            raise ServiceError("Файл повреждён или не удалось извлечь текст. Попробуйте сохранить его заново.", "unreadable_file", 422)
        if not text.strip():
            raise ServiceError("Читаемый текст не найден. Для фото нужны отчётливые маркировка и артикул.", "no_text", 422)
        rows, ignored = self.rows(text)
        if not rows:
            raise ServiceError("В файле не найдены строки спецификации.", "no_product_rows", 422)
        matched, unmatched, warnings = [], [], []
        started = time.monotonic()
        # Small batches bound the number of outstanding catalog requests.
        for offset in range(0, len(rows), 4):
            if time.monotonic() - started > 20:
                unmatched.extend({**row, "status": "Лимит времени обработки; загрузите отдельным файлом."} for row in rows[offset:])
                warnings.append("processing_budget_exceeded")
                break
            with ThreadPoolExecutor(max_workers=4) as pool:
                for hit, miss in pool.map(self._match_row, rows[offset:offset + 4]):
                    if hit:
                        matched.append(hit)
                    if miss:
                        unmatched.append(miss)
        total = round(sum(item["subtotal"] or 0 for item in matched), 2)
        complete = bool(matched) and not unmatched and all(item["subtotal"] is not None for item in matched)
        return {"total_positions_found": len(matched), "total_estimate_kzt": total,
                "estimate_complete": complete, "matched_items": matched, "unmatched_items": unmatched,
                "ignored_lines": ignored, "warnings": warnings, "cart_updated": False,
                "summary_text": f"Найдено позиций: {len(matched)}. Известная часть сметы: {total:,.2f} ₸. "
                                f"Несопоставленных строк: {len(unmatched)}. " +
                                ("Смета рассчитана по всем строкам." if complete else "Смета неполная: требуются уточнения.")}


spec_parser = SpecificationParser()
