import io
from unittest.mock import patch

import pymupdf
from docx import Document
from openpyxl import Workbook
from PIL import Image

from backend.tests.support import APIHarness
from spec_parser import MAX_BYTES, SpecificationParser


class UploadTests(APIHarness):
    def test_empty_leading_table_cells_preserve_quantity_column(self):
        text = "№\tАртикул\tКоличество\n\tTEST-1001\t3"
        response = self.upload("spec.txt", text.encode())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["estimate"]["matched_items"][0]["quantity"], 3)

    def test_total_request_limit_handles_chunked_stream(self):
        chunks = (b"x" * 1024 * 1024 for _ in range(17))
        response = self.client.post("/api/agent/upload-spec", content=chunks,
                                    headers={"Content-Type": "multipart/form-data; boundary=example",
                                             "Origin": "http://localhost:5173"})
        self.assertEqual(response.status_code, 413)
        self.assertEqual(response.json()["code"], "file_too_large")
        self.assertEqual(response.headers["access-control-allow-origin"], "http://localhost:5173")

    def upload(self, name, data, content_type="application/octet-stream"):
        return self.client.post("/api/agent/upload-spec", files={"file": (name, data, content_type)})

    def test_txt_quantities_headers_and_unknown_line(self):
        response = self.upload("spec.txt", "Спецификация\nTEST-1001 2 шт.\nUNKNOWN-999 3 шт.".encode())
        self.assertEqual(response.status_code, 200)
        estimate = response.json()["estimate"]
        self.assertEqual(estimate["matched_items"][0]["quantity"], 2)
        self.assertEqual(estimate["total_estimate_kzt"], 2000)
        self.assertEqual(len(estimate["unmatched_items"]), 1)
        self.assertFalse(estimate["estimate_complete"])
        self.assertEqual(self.cart()["total_items"], 0)

    def test_pdf_pipeline(self):
        with pymupdf.open() as document:
            page = document.new_page()
            page.insert_text((72, 72), "TEST-1001 2 pcs")
            data = document.tobytes()
        response = self.upload("spec.pdf", data, "application/pdf")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["estimate"]["matched_items"][0]["quantity"], 2)

    def test_word_table_pipeline(self):
        document = Document()
        table = document.add_table(rows=2, cols=3)
        for cell, value in zip(table.rows[0].cells, ["Артикул", "Количество", "Ед.изм."]):
            cell.text = value
        for cell, value in zip(table.rows[1].cells, ["TEST-1001", "2", "шт."]):
            cell.text = value
        buffer = io.BytesIO()
        document.save(buffer)
        response = self.upload("spec.docx", buffer.getvalue())
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["estimate"]["matched_items"][0]["quantity"], 2)

    def test_excel_quantity_column_is_not_assumed_last(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Количество", "Артикул", "Наименование", "Цена"])
        sheet.append([3.0, "TEST-1001", "Автомат", 50000])
        buffer = io.BytesIO()
        workbook.save(buffer)
        response = self.upload("spec.xlsx", buffer.getvalue())
        self.assertEqual(response.status_code, 200, response.text)
        row = response.json()["estimate"]["matched_items"][0]
        self.assertEqual(row["quantity"], 3)
        self.assertEqual(row["unit_price"], 1000)

    def test_missing_quantity_does_not_become_one(self):
        response = self.upload("spec.txt", b"TEST-1001")
        estimate = response.json()["estimate"]
        self.assertIsNone(estimate["matched_items"][0]["quantity"])
        self.assertIsNone(estimate["matched_items"][0]["subtotal"])
        self.assertFalse(estimate["estimate_complete"])

    def test_ambiguous_names_remain_unresolved(self):
        parser = SpecificationParser(self.catalog)
        with patch.object(self.catalog, "search_products", return_value=[
            {"id": 1001, "name": "A", "_match_score": 3, "_match_coverage": 1},
            {"id": 1002, "name": "B", "_match_score": 3, "_match_coverage": 1},
        ]):
            result = parser.parse_specification("Автомат 16А 2 шт.".encode(), "test.txt")
        self.assertEqual(result["matched_items"], [])
        self.assertEqual(len(result["unmatched_items"][0]["candidates"]), 2)

    def test_validation_of_empty_oversized_wrong_signature_and_legacy_word(self):
        for name, content, status in (("spec.txt", b"", 422), ("spec.exe", b"bad", 415),
                                      ("spec.pdf", b"not pdf", 415), ("spec.docx", b"not zip", 415),
                                      ("spec.doc", b"old word", 415), ("spec.txt", b"x" * (MAX_BYTES + 1), 413)):
            with self.subTest(name=name, status=status):
                self.assertEqual(self.upload(name, content).status_code, status)

    def test_image_ocr_and_missing_runtime(self):
        buffer = io.BytesIO()
        Image.new("RGB", (300, 100), "white").save(buffer, format="PNG")
        with patch.object(SpecificationParser, "_ocr", return_value="TEST-1001 2 pcs"):
            response = self.upload("spec.png", buffer.getvalue())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["estimate"]["matched_items"][0]["quantity"], 2)
        with patch("spec_parser.ocr_available", return_value=False):
            response = self.upload("spec.png", buffer.getvalue())
        self.assertEqual(response.status_code, 503)
        self.assertEqual(response.json()["code"], "ocr_unavailable")
